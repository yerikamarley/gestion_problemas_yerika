import unittest
from io import BytesIO
from unittest.mock import MagicMock, patch

import pandas as pd
from openpyxl import load_workbook

import app_logic
from components.seguimiento_proveedores import exportar_csv, tabla_objetivos
from core.permissions import ROLE_ADMIN, ROLE_SOPORTE, VIEW_SEGUIMIENTO_PROVEEDORES
from repositories.proveedores import leer_candidatos_proveedores
from services.proveedores import (
    CONFIRMADO, MENCION, SIN_CLIENTE, SIN_PROVEEDOR, catalogo_ans,
    evaluar_ans_proveedores, festivos_colombia, filtrar_periodo_proveedores,
    minutos_servicio, plazo_dias_habiles, prioridad_ans, registros_proveedores,
    resumen_ans_proveedores,
)
from utils.proveedores_excel_export import COLUMNAS_INFORME, exportar_informe_proveedores, informe_proveedores


def fecha(valor):
    return pd.Timestamp(valor, tz="America/Bogota")


def caso(**cambios):
    return {"numero": "CS1", "cuenta": "Cliente Uno", "estado": "Nuevo",
            "prioridad": "3 - Moderada", "creado": "2026-10-09 08:00:00", **cambios}


def analizar(row, corte="2026-10-10 12:00:00"):
    df = registros_proveedores(pd.DataFrame([caso(**row)]), pd.DataFrame())
    return evaluar_ans_proveedores(df, fecha(corte))


class DeteccionProveedoresTest(unittest.TestCase):
    def test_nombres_en_notas_e_historial_y_clientes_en_ambos_tipos(self):
        casos = pd.DataFrame([caso(notas_trabajo="Rpost revisa la entrega"), caso(numero="CS2", historial_conversacion=[
            {"tipo": "Nota de Trabajo", "contenido": "Se escaló a Autentic para revisar la caída"}])])
        incidentes = pd.DataFrame([{"numero": "INC1", "empresa": "Cliente Dos", "descripcion": "Software Colombia presenta falla", "creado": "2026-10-09"}])
        df = registros_proveedores(casos, incidentes)
        self.assertEqual(3, len(df))
        self.assertEqual(MENCION, df.iloc[0]["Relación"])
        self.assertEqual(CONFIRMADO, df.iloc[1]["Relación"])
        self.assertEqual("Cliente Dos", df.iloc[2]["Cliente"])

    def test_certicamara_y_firma_generica_no_son_proveedor(self):
        df = analizar({"descripcion": "Certicámara: necesito instalar token físico y firma digital"})
        self.assertTrue(df.empty)

    def test_footer_y_correos_citados_no_son_evidencia(self):
        df = analizar({"historial_conversacion": [{"tipo": "Correo entrante", "contenido": "Falla de acceso\nAviso de Confidencialidad: Rpost Autentic"}]})
        self.assertTrue(df.empty)

    def test_solicitud_y_negacion_no_confirman_escalamiento(self):
        for texto in ["No se escaló a Rpost", "Pendiente de escalamiento a Rpost", "Solicitamos escalar a Rpost", "Por escalar a Rpost"]:
            with self.subTest(texto=texto):
                df = analizar({"notas_trabajo": texto})
                self.assertEqual(MENCION, df.iloc[0]["Relación"])

    def test_mencion_cliente_no_confirma_escalamiento(self):
        df = analizar({"descripcion": "Caso escalado a Rpost según el cliente"})
        self.assertEqual(MENCION, df.iloc[0]["Relación"])

    def test_flag_y_nombre_ganan_sin_asignar_otro_mencionado(self):
        df = analizar({"nombre_proveedor": "Rpost", "escalado_proveedor": "sí", "notas_trabajo": "Consultar luego con Autentic"})
        relaciones = dict(zip(df["Proveedor"], df["Relación"]))
        self.assertEqual(CONFIRMADO, relaciones["RPost / eSoluciona"])
        self.assertEqual(MENCION, relaciones["AutenTIC"])

    def test_escalamiento_no_se_extiende_a_mencion_secundaria_en_misma_frase(self):
        df = analizar({"notas_trabajo": "Se escaló a Rpost y se considera consultar a Autentic"})
        relaciones = dict(zip(df["Proveedor"], df["Relación"]))
        self.assertEqual(CONFIRMADO, relaciones["RPost / eSoluciona"])
        self.assertEqual(MENCION, relaciones["AutenTIC"])

    def test_escalamiento_conjunto_confirmado_ambos_sin_duplicar_registro(self):
        df = analizar({"notas_trabajo": "Se escaló a Rpost y Autentic"})
        self.assertEqual(2, len(df))
        self.assertTrue((df["Relación"] == CONFIRMADO).all())
        self.assertEqual(1, len(df.drop_duplicates(["Tipo", "Número"])))

    def test_flag_sin_nombre_preserva_pendiente_de_identificacion(self):
        df = analizar({"escalado_proveedor": True})
        self.assertEqual(SIN_PROVEEDOR, df.iloc[0]["Proveedor"])
        self.assertEqual("Sin ANS definido", df.iloc[0]["ANS Solución"])

    def test_cliente_no_se_inventa_a_partir_del_footer(self):
        df = analizar({"cuenta": None, "contacto": "Persona", "notas_trabajo": "Se escaló a Rpost"})
        self.assertEqual(SIN_CLIENTE, df.iloc[0]["Cliente"])

    def test_cuenta_vacia_de_exportacion_no_cuenta_como_cliente(self):
        df = analizar({"cuenta": "(vacío)", "nombre_proveedor": "Rpost"})
        self.assertEqual(SIN_CLIENTE, df.iloc[0]["Cliente"])
        self.assertEqual(0, resumen_ans_proveedores(df).iloc[0]["Clientes"])

    def test_varios_proveedores_no_mezclan_fechas(self):
        df = analizar({"notas_trabajo": "Proveedor: Rpost\nEscalamiento al proveedor: 09/10/2026 08:00\nPrimera respuesta del proveedor: 09/10/2026 08:20\n"
                                            "Proveedor: Autentic\nEscalamiento al proveedor: 09/10/2026 10:00\nSolución del proveedor: 09/10/2026 11:00"})
        rpost = df[df["Proveedor"] == "RPost / eSoluciona"].iloc[0]
        autentic = df[df["Proveedor"] == "AutenTIC"].iloc[0]
        self.assertEqual(fecha("2026-10-09 08:00"), rpost["Escalamiento"])
        self.assertTrue(pd.isna(rpost["Solución"]))
        self.assertEqual(fecha("2026-10-09 11:00"), autentic["Solución"])
        self.assertTrue(pd.isna(autentic["Primera respuesta"]))

    def test_no_asigna_fecha_ambigua_a_dos_proveedores(self):
        df = analizar({"notas_trabajo": "Rpost y Autentic: Escalamiento al proveedor: 09/10/2026 08:00"})
        self.assertTrue(df["Escalamiento"].isna().all())

    def test_historial_incidente_y_fecha_encabezado_no_sustituyen_hitos(self):
        incidentes = pd.DataFrame([{"numero": "INC1", "observaciones_trabajo":
            "2026-10-09 13:00:00 | Nota de Trabajo | soporte | | Se escaló a Rpost", "prioridad": "P1"}])
        df = registros_proveedores(pd.DataFrame(), incidentes)
        self.assertEqual(CONFIRMADO, df.iloc[0]["Relación"])
        self.assertTrue(pd.isna(df.iloc[0]["Escalamiento"]))


class CalendarioANSTest(unittest.TestCase):
    def test_catalogo_corresponde_a_matriz_del_excel(self):
        cat = catalogo_ans()["proveedores"]
        self.assertEqual(4, len(cat))
        self.assertEqual(600, cat["Software Colombia"]["prioridades"]["4"]["objetivos"]["solucion"])
        self.assertIsNone(cat["AutenTIC"]["prioridades"]["1"]["objetivos"]["diagnostico"])
        self.assertEqual(45, cat["RPost / eSoluciona"]["prioridades"]["2"]["objetivos"]["atencion"])
        self.assertEqual(4, len(tabla_objetivos("AutenTIC")))

    def test_prioridad_desconocida_y_p5_no_reciben_ans_p4(self):
        for valor in [None, "", "5 - Muy baja", "P5", "sin prioridad"]:
            self.assertIsNone(prioridad_ans(valor))
        self.assertEqual([1, 2, 3, 4], [prioridad_ans(v) for v in ["P1", "2 - Alta", "Moderada", "4 - Baja"]])

    def test_festivos_fijos_trasladados_y_semana_santa(self):
        festivos = festivos_colombia(2026)
        for f in ["2026-01-01", "2026-01-12", "2026-04-02", "2026-04-03", "2026-05-18", "2026-06-08", "2026-06-15", "2026-10-12"]:
            self.assertIn(pd.Timestamp(f).date(), festivos)
        self.assertNotIn(pd.Timestamp("2026-10-09").date(), festivos)

    def test_rpost_viernes_sabado_domingo_y_lunes_festivo(self):
        tiempo = minutos_servicio(fecha("2026-10-09 17:00"), fecha("2026-10-13 08:00"), "RPost / eSoluciona", 1)
        self.assertEqual(1080, tiempo)  # 60 viernes + 300 sábado + 660 lunes + 60 martes; festivos sin excepción en el Excel.

    def test_autentic_critico_continuo_otras_prioridades_horario_festivo(self):
        a, b = fecha("2026-10-11 18:00"), fecha("2026-10-12 09:00")
        self.assertEqual(900, minutos_servicio(a, b, "AutenTIC", 1))
        self.assertEqual(120, minutos_servicio(a, b, "AutenTIC", 2))

    def test_grafologos_menor_ventana_diaria_y_mayor_24h(self):
        a, b = fecha("2026-10-09 21:00"), fecha("2026-10-10 09:00")
        self.assertEqual(720, minutos_servicio(a, b, "Grafólogos Forenses", 1))
        self.assertEqual(120, minutos_servicio(a, b, "Grafólogos Forenses", 3))

    def test_tres_dias_habiles_no_son_72_horas(self):
        self.assertEqual(fecha("2026-10-15 10:00"), plazo_dias_habiles(fecha("2026-10-09 10:00"), 3))


class EvaluacionANSTest(unittest.TestCase):
    def test_un_segundo_fuera_del_objetivo_no_se_redondea_a_cumple(self):
        df = analizar({"prioridad": "P1", "notas_trabajo": "Rpost\nEscalamiento: 09/10/2026 08:00:00\nSolución: 09/10/2026 09:00:01"})
        self.assertEqual("No cumple", df.iloc[0]["ANS Solución"])

    def test_mide_desde_escalamiento_en_vez_de_creacion(self):
        df = analizar({"prioridad": "P1", "notas_trabajo": "Proveedor: Rpost\nEscalamiento al proveedor: 09/10/2026 10:00\n"
                      "Primera respuesta del proveedor: 09/10/2026 10:30\nSolución del proveedor: 09/10/2026 11:00"})
        self.assertEqual("Cumple", df.iloc[0]["ANS Atención"])
        self.assertEqual("Cumple", df.iloc[0]["ANS Solución"])
        self.assertEqual(60, df.iloc[0]["Minutos Solución"])

    def test_abiertos_en_plazo_y_vencidos_separados_de_evaluados(self):
        df = analizar({"prioridad": "P1", "notas_trabajo": "Rpost\nEscalamiento: 09/10/2026 10:00"}, "2026-10-09 10:40")
        self.assertEqual("Vencido", df.iloc[0]["ANS Atención"])
        self.assertEqual("En plazo", df.iloc[0]["ANS Solución"])
        resumen = resumen_ans_proveedores(df).iloc[0]
        self.assertEqual(0, resumen["Atención evaluados"])
        self.assertTrue(pd.isna(resumen["Atención % cumplimiento"]))
        self.assertEqual(1, resumen["Atención vencidos"])

    def test_cerrado_sin_fecha_proveedor_es_sin_datos(self):
        df = analizar({"estado": "Cerrado", "cerrado": "2026-10-09 12:00", "prioridad": "P1",
                       "notas_trabajo": "Rpost\nEscalamiento: 09/10/2026 10:00"})
        self.assertEqual("Sin datos", df.iloc[0]["ANS Solución"])

    def test_no_inventa_fechas_de_historial(self):
        df = analizar({"escalado_proveedor": "true", "nombre_proveedor": "Autentic", "cerrado": "2026-10-09 12:00"})
        self.assertEqual("Sin datos", df.iloc[0]["ANS Solución"])

    def test_fecha_anterior_o_invalida_no_cumple_falsamente(self):
        for fecha_respuesta in ["09/10/2026 07:00", "31/02/2026 10:00"]:
            df = analizar({"notas_trabajo": "Rpost\nEscalamiento: 09/10/2026 10:00\nPrimera respuesta del proveedor: " + fecha_respuesta})
            self.assertEqual("Fechas inconsistentes", df.iloc[0]["ANS Atención"])

    def test_dos_ciclos_mismo_proveedor_conservan_sus_soluciones(self):
        df = analizar({"notas_trabajo": "Rpost\nEscalamiento: 09/09/2026 10:00\nSolución del proveedor: 09/09/2026 11:00\n"
                       "Escalamiento: 09/10/2026 10:00\nSolución del proveedor: 09/10/2026 13:00"})
        self.assertEqual(2, len(df))
        self.assertEqual(fecha("2026-09-09 11:00"), df.iloc[0]["Solución"])
        self.assertEqual(fecha("2026-10-09 13:00"), df.iloc[1]["Solución"])
        mensual = filtrar_periodo_proveedores(df, fecha("2026-10-01"), fecha("2026-11-01"))
        self.assertEqual(1, len(mensual))

    def test_historico_no_usa_solucion_posterior_al_corte(self):
        df = analizar({"prioridad": "P1", "notas_trabajo": "Software Colombia\nEscalamiento: 30/09/2026 22:00\nSolución del proveedor: 01/10/2026 03:00"}, "2026-09-30 23:59:59")
        self.assertEqual("En plazo", df.iloc[0]["ANS Solución"])

    def test_mes_incluye_creado_anterior_escalado_ahora_y_pendientes(self):
        df = registros_proveedores(pd.DataFrame([
            caso(numero="CS1", creado="2026-08-01", notas_trabajo="Autentic\nEscalamiento: 03/10/2026 08:00"),
            caso(numero="CS2", creado="2026-08-01", notas_trabajo="Rpost\nEscalamiento: 01/09/2026 08:00"),
            caso(numero="CS3", creado="2026-08-01", notas_trabajo="Rpost\nEscalamiento: 01/09/2026 08:00\nSolución: 01/09/2026 09:00"),
        ]), pd.DataFrame())
        mensual = filtrar_periodo_proveedores(df, fecha("2026-10-01"), fecha("2026-11-01"))
        self.assertEqual({"CS1", "CS2"}, set(mensual["Número"]))

    def test_resumen_denomina_evaluados_y_preserva_menciones(self):
        df = registros_proveedores(pd.DataFrame([
            caso(numero="CS1", prioridad="P1", notas_trabajo="Rpost\nEscalamiento: 09/10/2026 08:00\nSolución: 09/10/2026 09:00"),
            caso(numero="CS2", prioridad="P1", notas_trabajo="Rpost\nEscalamiento: 09/10/2026 08:00\nSolución: 09/10/2026 09:01"),
            caso(numero="CS3", descripcion="Rpost", cuenta="Cliente Dos"),
        ]), pd.DataFrame())
        resumen = resumen_ans_proveedores(evaluar_ans_proveedores(df, fecha("2026-10-10 12:00"))).iloc[0]
        self.assertEqual(50, resumen["Solución % cumplimiento"])
        self.assertEqual(2, resumen["Solución evaluados"])
        self.assertEqual(1, resumen["Menciones"])
        self.assertEqual(2, resumen["Clientes"])


class ConsultaYExportacionTest(unittest.TestCase):
    def test_informe_conserva_historial_estado_y_un_registro_por_proveedor_sin_duplicar_ciclos(self):
        casos = pd.DataFrame([caso(estado="Resuelto", notas_trabajo="Rpost\nEscalamiento: 08/10/2026 08:00\n"
                                  "Solución: 08/10/2026 09:00\nEscalamiento: 09/10/2026 08:00",
                                  historial_conversacion=[{"fecha": "2026-10-09 10:00", "tipo": "Nota de Trabajo",
                                                          "usuario": "soporte", "contenido": "Proveedor confirma revisión completa"}],
                                  actualizaciones="Actualización posterior sin nombre del proveedor")])
        incidentes = pd.DataFrame([{"numero": "INC1", "estado": "Activo", "creado": "2026-10-09 13:00:00Z",
                                   "nombre_proveedor": "Autentic", "observaciones_trabajo": "Seguimiento completo del incidente"}])
        registros = registros_proveedores(casos, incidentes)
        self.assertEqual(3, len(registros))
        reporte = informe_proveedores(registros)
        self.assertEqual(COLUMNAS_INFORME, reporte.columns.tolist())
        self.assertEqual(2, len(reporte))
        historial = reporte.loc[reporte["Caso o incidente"] == "Caso CS1", "Historial del caso o incidente"].iloc[0]
        self.assertIn("Proveedor confirma revisión completa", historial)
        self.assertIn("2026-10-09 10:00 · Nota de Trabajo · soporte", historial)
        self.assertIn("Actualización posterior sin nombre del proveedor", historial)
        libro = load_workbook(BytesIO(exportar_informe_proveedores(registros)))
        hoja = libro.active
        self.assertEqual("A2", hoja.freeze_panes)
        self.assertEqual(1, len(hoja.tables))
        self.assertEqual("Incidente INC1", hoja["C2"].value)
        self.assertEqual("Activo", hoja["E2"].value)
        self.assertEqual(8, hoja["B2"].value.hour)
        self.assertIsNone(hoja["B2"].value.tzinfo)
        self.assertTrue(hoja["B2"].is_date)

    def test_excel_no_ejecuta_formulas_ni_trunca_historiales_largos(self):
        historial = "=HYPERLINK(x)\n" + "Seguimiento detallado\n" * 4000
        datos = registros_proveedores(pd.DataFrame([caso(nombre_proveedor="Rpost", notas_trabajo=historial)]), pd.DataFrame())
        texto_original = informe_proveedores(datos).iloc[0]["Historial del caso o incidente"]
        datos.loc[:, "Proveedor"] = "=SUM(1,2)"
        hoja = load_workbook(BytesIO(exportar_informe_proveedores(datos))).active
        self.assertGreater(hoja.max_row, 2)
        self.assertEqual(texto_original, "".join(hoja.cell(fila, 4).value for fila in range(2, hoja.max_row + 1)))
        for fila in hoja.iter_rows(min_row=2):
            self.assertEqual("=SUM(1,2)", fila[0].value)
            self.assertEqual("s", fila[0].data_type)
            self.assertFalse(any(celda.data_type == "f" for celda in fila))

    def test_pantalla_carga_resumen_filtros_e_informe_sin_excepciones(self):
        from streamlit.testing.v1 import AppTest
        script = '''import pandas as pd
from components.seguimiento_proveedores import render_seguimiento_proveedores
from services.proveedores import registros_proveedores
ahora = pd.Timestamp.now(tz="America/Bogota").strftime("%Y-%m-%d")
casos = pd.DataFrame([{"numero":"CS1","cuenta":"Cliente Uno","estado":"Nuevo","creado":ahora,"prioridad":"P1","nombre_proveedor":"Rpost","escalado_proveedor":"si"}])
render_seguimiento_proveedores(lambda: registros_proveedores(casos, pd.DataFrame()))
'''
        app = AppTest.from_string(script).run(timeout=20)
        self.assertFalse(app.exception)
        self.assertEqual("1", app.metric[0].value)
        self.assertEqual(6, len(app.dataframe))
        self.assertTrue(any(boton.proto.label == "Descargar informe de proveedores (Excel)" for boton in app.get("download_button")))
        app.selectbox(key="proveedores_filtro").select("AutenTIC").run(timeout=20)
        self.assertFalse(app.exception)
        self.assertTrue(app.info)

    def test_backend_rechaza_sin_vista_autorizada_antes_de_consultar(self):
        with patch.object(app_logic, "exigir_contexto_lectura", side_effect=app_logic.AutorizacionError), patch("repositories.proveedores.leer_candidatos_proveedores") as lectura:
            with self.assertRaises(app_logic.AutorizacionError):
                app_logic.load_seguimiento_proveedores()
            lectura.assert_not_called()

    def test_consulta_candidatos_no_limita_creacion_ni_pagina_y_cierra(self):
        conexion = MagicMock()
        cursor = MagicMock()
        cursor.fetchall.return_value = [("CS1", "Rpost", "true")]
        with patch("repositories.proveedores.exigir_contexto_consulta"), patch("repositories.proveedores.get_conn", return_value=conexion), patch("repositories.proveedores.columnas_existentes", return_value={"numero", "nombre_proveedor", "escalado_proveedor"}), patch("repositories.proveedores.db_execute", return_value=cursor) as ejecutar:
            df = leer_candidatos_proveedores("cases", ["numero", "nombre_proveedor", "escalado_proveedor"])
        sql, params = ejecutar.call_args.args[1:]
        self.assertNotIn("LIMIT", sql)
        self.assertNotIn("creado >=", sql)
        self.assertIn("ILIKE ANY", sql)
        self.assertIn("%rpost%", params)
        self.assertEqual("CS1", df.iloc[0]["numero"])
        conexion.close.assert_called_once()

    def test_exporta_fechas_colombia_y_protege_formulas(self):
        csv = exportar_csv(pd.DataFrame({"Nota": ["=HYPERLINK(x)"], "Fecha": [fecha("2026-10-09 08:00")]})).decode("utf-8-sig")
        self.assertIn("'=HYPERLINK(x)", csv)
        self.assertIn("09/10/2026 08:00:00", csv)
        self.assertNotIn("-05:00", csv)

    def test_una_vista_y_rutas_anteriores_se_resuelven_sin_ampliar_acceso(self):
        import app_ui
        catalogo = [v for v in app_ui.VIEW_CATALOG if "proveedor" in v[0]]
        self.assertEqual([VIEW_SEGUIMIENTO_PROVEEDORES], [v[0] for v in catalogo])
        for vieja in ["seguimiento_rpost", "seguimiento_autentic"]:
            self.assertEqual((VIEW_SEGUIMIENTO_PROVEEDORES, False), app_ui.resolver_vista_permitida(ROLE_ADMIN, vieja))
            self.assertTrue(app_ui.resolver_vista_permitida(ROLE_SOPORTE, vieja)[1])


if __name__ == "__main__":
    unittest.main()
