"""Informe operativo de proveedores con fechas nativas e historial completo."""

from io import BytesIO
from math import ceil

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


COLUMNAS_INFORME = ["Nombre del proveedor", "Fecha", "Caso o incidente",
                    "Historial del caso o incidente", "Estado del caso"]


def informe_proveedores(datos):
    filas = []
    # Los ciclos de escalamiento no duplican el mismo caso en el informe.
    for row in datos.drop_duplicates(["Proveedor", "Tipo", "Número"], keep="last").to_dict("records"):
        creado = row.get("Creado")
        fecha = pd.Timestamp(creado) if pd.notna(creado) else pd.NaT
        if pd.notna(fecha) and fecha.tzinfo is not None:
            fecha = fecha.tz_convert("America/Bogota").tz_localize(None)
        historial = row.get("Historial")
        historial = historial if isinstance(historial, str) and historial.strip() else "Sin historial disponible"
        filas.append({"Nombre del proveedor": row["Proveedor"], "Fecha": fecha,
                      "Caso o incidente": f"{row['Tipo']} {row['Número']}",
                      "Historial del caso o incidente": historial, "Estado del caso": row.get("Estado", "")})
    return pd.DataFrame(filas, columns=COLUMNAS_INFORME).sort_values(
        ["Nombre del proveedor", "Fecha", "Caso o incidente"], na_position="last", kind="stable").reset_index(drop=True)


def exportar_informe_proveedores(datos):
    reporte = informe_proveedores(datos)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Informe proveedores"
    hoja.append(COLUMNAS_INFORME)
    for celda in hoja[1]:
        celda.fill = PatternFill("solid", fgColor="5D16A6")
        celda.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    hoja.row_dimensions[1].height = 32
    for fila in reporte.itertuples(index=False, name=None):
        proveedor, fecha, numero, historial, estado = fila
        historial = ILLEGAL_CHARACTERS_RE.sub("", historial)
        # Excel admite 32 767 caracteres por celda. Continuar en otra fila
        # conserva todo el historial en vez de truncarlo silenciosamente.
        for comienzo in range(0, max(1, len(historial)), 30000):
            parte = historial[comienzo:comienzo + 30000]
            valores = [proveedor, fecha.to_pydatetime() if pd.notna(fecha) else None, numero, parte, estado]
            indice = hoja.max_row + 1
            for columna, valor in enumerate(valores, 1):
                celda = hoja.cell(indice, columna)
                if isinstance(valor, str):
                    celda.value = ILLEGAL_CHARACTERS_RE.sub("", valor)
                    celda.data_type = "s"  # El texto de origen nunca se ejecuta como fórmula.
                else:
                    celda.value = valor
                celda.font = Font(name="Calibri", size=11)
                celda.alignment = Alignment(vertical="top", wrap_text=True)
            hoja.cell(indice, 2).number_format = "dd/mm/yyyy hh:mm"
            lineas = sum(max(1, ceil(len(linea) / 95)) for linea in parte.splitlines())
            hoja.row_dimensions[indice].height = min(210, max(48, lineas * 15))
    for columna, ancho in {"A": 30, "B": 22, "C": 25, "D": 100, "E": 23}.items():
        hoja.column_dimensions[columna].width = ancho
    if hoja.max_row > 1:
        tabla = Table(displayName="InformeProveedores", ref=f"A1:E{hoja.max_row}")
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        hoja.add_table(tabla)
    hoja.freeze_panes = "A2"
    hoja.sheet_view.showGridLines = False
    hoja.sheet_view.zoomScale = 90
    hoja.print_title_rows = "1:1"
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A3
    hoja.page_setup.fitToWidth, hoja.page_setup.fitToHeight = 1, 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
